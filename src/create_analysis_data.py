"""Create analysis-ready data from audited Working_Data without changing raw data."""
from __future__ import annotations

import csv
import hashlib
import json
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "RAW_DATA.csv"
MAIN = ROOT / "KKU_Financial_Cleaning.xlsx"
DICTIONARY = ROOT / "Data_Dictionary.xlsx"
LOGBOOK = ROOT / "Cleaning_Log.xlsx"
SOURCE_HASH = "a4b907893c14ae9b57a282200614f6753e8e7aefe484f957810641f955534e4d"
INCOME = ["FamilyIncome_Clean", "PartTimeIncome_Clean", "ScholarshipIncome_Clean", "OtherIncome_Clean"]
EXPENSE = ["FoodExpense_Clean", "TransportExpense_Clean", "HousingExpense_Clean", "UtilityExpense_Clean", "EducationExpense_Clean", "EntertainmentExpense_Clean", "ShoppingExpense_Clean", "SubscriptionExpense_Clean", "OtherExpense_Clean"]
RATIOS = ["FoodRatio", "TransportRatio", "HousingRatio", "UtilityRatio", "EducationRatio", "EntertainmentRatio", "ShoppingRatio", "SubscriptionRatio", "OtherExpenseRatio"]
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def sheet_matrix(ws):
    return [[cell.value for cell in row] for row in ws.iter_rows()]


def style_data_sheet(ws):
    ws.freeze_panes = "A2"; ws.sheet_view.showGridLines = False
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.auto_filter.ref = ws.dimensions
    for column in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(column)].width = min(max(14, len(str(ws.cell(1, column).value)) + 2), 28)


def is_present(value):
    return value is not None and value != ""


def clean_number(value):
    return None if not is_present(value) else float(value)


def main():
    assert sha256(RAW) == SOURCE_HASH, "Source hash differs; stop before modifying outputs."
    with RAW.open(encoding="utf-8-sig", newline="") as f:
        raw_matrix = list(csv.reader(f))

    wb = load_workbook(MAIN)
    raw_before = sheet_matrix(wb["Raw_Copy"])
    work_before = sheet_matrix(wb["Working_Data"])
    assert [["" if x is None else str(x) for x in row] for row in raw_before] == raw_matrix, "Raw_Copy no longer matches source."
    work_headers = work_before[0]
    work_records = [dict(zip(work_headers, row)) for row in work_before[1:]]
    assert len(work_records) == 30

    headers = [
        "Respondent_ID", "Year", "Faculty", "Accommodation", "PartTime",
        "FamilyIncome", "PartTimeIncome", "ScholarshipIncome", "OtherIncome", "TotalIncome",
        "FoodExpense", "TransportExpense", "HousingExpense", "UtilityExpense", "EducationExpense",
        "EntertainmentExpense", "ShoppingExpense", "SubscriptionExpense", "OtherExpense", "TotalExpense",
        "NetBalance", "ExpenseRatio", *RATIOS,
        "SavingStatus", "SavingAmount", "SavingPattern",
        "PlanSpending_Code", "TrackExpenses_Code", "ControlExpenses_Code", "PrioritizeNeeds_Code", "ImpulsePurchase_Code",
        "MoneySufficiency_Code", "MoneySufficiency_Group", "ShortageMonths_Code",
        "IncomeComplete", "ExpenseComplete", "FinancialIndicatorsAvailable",
    ]
    rows = []
    for record in work_records:
        incomes = [clean_number(record[x]) for x in INCOME]
        expenses = [clean_number(record[x]) for x in EXPENSE]
        income_complete = all(x is not None for x in incomes)
        expense_complete = all(x is not None for x in expenses)
        total_income = sum(incomes) if income_complete else None
        total_expense = sum(expenses) if expense_complete else None
        net_balance = total_income - total_expense if total_income is not None and total_expense is not None else None
        expense_ratio = total_expense / total_income * 100 if total_income not in (None, 0) and total_expense is not None else None
        ratios = [expense / total_expense * 100 if total_expense not in (None, 0) and expense is not None else None for expense in expenses]
        output = {
            "Respondent_ID": record["Row_ID"], "Year": record["Year"], "Faculty": record["Faculty"],
            "Accommodation": record["Accommodation"], "PartTime": record["PartTime"],
            "FamilyIncome": incomes[0], "PartTimeIncome": incomes[1], "ScholarshipIncome": incomes[2], "OtherIncome": incomes[3], "TotalIncome": total_income,
            "FoodExpense": expenses[0], "TransportExpense": expenses[1], "HousingExpense": expenses[2], "UtilityExpense": expenses[3], "EducationExpense": expenses[4],
            "EntertainmentExpense": expenses[5], "ShoppingExpense": expenses[6], "SubscriptionExpense": expenses[7], "OtherExpense": expenses[8], "TotalExpense": total_expense,
            "NetBalance": net_balance, "ExpenseRatio": expense_ratio,
            "SavingStatus": record["SavingStatus"], "SavingAmount": clean_number(record["SavingAmount_Clean"]), "SavingPattern": record["SavingPattern"],
            "PlanSpending_Code": record["PlanSpending_Code"], "TrackExpenses_Code": record["TrackExpenses_Code"], "ControlExpenses_Code": record["ControlExpenses_Code"],
            "PrioritizeNeeds_Code": record["PrioritizeNeeds_Code"], "ImpulsePurchase_Code": record["ImpulsePurchase_Code"],
            "MoneySufficiency_Code": record["MoneySufficiency_Code"], "MoneySufficiency_Group": record["MoneySufficiency_Group"], "ShortageMonths_Code": record["ShortageMonths_Code"],
            "IncomeComplete": int(income_complete), "ExpenseComplete": int(expense_complete), "FinancialIndicatorsAvailable": int(total_income is not None and total_expense is not None),
        }
        output.update(dict(zip(RATIOS, ratios)))
        rows.append(output)

    if "Analysis_Data" in wb.sheetnames:
        del wb["Analysis_Data"]
    ws = wb.create_sheet("Analysis_Data")
    ws.append(headers)
    for record in rows:
        ws.append([record[x] if not (isinstance(record[x], float) and np.isnan(record[x])) else None for x in headers])
    style_data_sheet(ws)
    money_cols = [headers.index(x) + 1 for x in headers if x.endswith("Income") or x.endswith("Expense") or x in {"TotalIncome", "TotalExpense", "NetBalance", "SavingAmount"}]
    ratio_cols = [headers.index(x) + 1 for x in ["ExpenseRatio", *RATIOS]]
    for col in money_cols:
        for row in range(2, ws.max_row + 1): ws.cell(row, col).number_format = '#,##0.00'
    for col in ratio_cols:
        for row in range(2, ws.max_row + 1): ws.cell(row, col).number_format = '0.00'

    # Document approved retention in the workbook report and audit log; no calculation is logged.
    quality = wb["Data_Quality_Report"]
    quality.append([]); quality.append(["Approved Review Resolution"])
    quality.append(["Item", "Decision", "Evidence", "Status"])
    quality.append(["HousingExpense = 32000", "Retained as reported", "RAW_DATA.csv contains exactly 32000; no contrary documentary evidence.", "POTENTIAL OUTLIER / RETAINED"])
    quality.append(["Likert direction", "No reversal applied", "CSV supplies 1–5 values but not response-anchor labels.", "LIMITATION DOCUMENTED"])

    log_ws = wb["Cleaning_Log"]
    log_headers = [cell.value for cell in log_ws[1]]
    existing = [dict(zip(log_headers, [c.value for c in row])) for row in log_ws.iter_rows(min_row=2)]
    existing = [x for x in existing if not (x["Variable"] == "HousingExpense" and x["Action"] == "Reviewed and retained as reported")]
    # Rewrite log only to make the single review record idempotent.
    log_ws.delete_rows(2, log_ws.max_row)
    for i, entry in enumerate(existing, 1):
        log_ws.append([i] + [entry.get(c, "") for c in log_headers[1:]])
    log_ws.append([len(existing)+1, 7, "HousingExpense", "32000", 32000, "possible_outlier", "Reviewed and retained as reported", "RAW_DATA.csv contains exactly 32000; no documentary evidence supports correction.", "ACCEPTED AS VALID", "Influential/potential outlier retained for sensitivity-aware future review."])
    style_data_sheet(log_ws)

    # Confirm protected sheets have no value changes before save.
    assert sheet_matrix(wb["Raw_Copy"]) == raw_before
    assert sheet_matrix(wb["Working_Data"]) == work_before
    wb.save(MAIN)

    # Synchronize standalone log with the canonical workbook log.
    lb = load_workbook(LOGBOOK)
    lws = lb["Cleaning_Log"]
    lws.delete_rows(1, lws.max_row)
    for row in sheet_matrix(load_workbook(MAIN)["Cleaning_Log"]): lws.append(row)
    style_data_sheet(lws)
    lb.save(LOGBOOK)

    # Extend the data dictionary without removing its existing entries.
    db = load_workbook(DICTIONARY)
    dws = db["Data_Dictionary"]
    dh = [cell.value for cell in dws[1]]
    entries = {row[0].value: row[0].row for row in dws.iter_rows(min_row=2) if row[0].value}
    dictionary_rows = {
        "TotalIncome": ["TotalIncome", "Sum of four clean income components", "FamilyIncome_Clean + PartTimeIncome_Clean + ScholarshipIncome_Clean + OtherIncome_Clean", "Numeric", "Ratio", "Derived financial indicator", "NA unless all four components are non-missing", "Strict missing propagation: do not treat NA as zero.", "Income level and later financial indicators", "Formula: sum of four components only when IncomeComplete=1."],
        "TotalExpense": ["TotalExpense", "Sum of nine clean expense components", "Nine cleaned expense variables", "Numeric", "Ratio", "Derived financial indicator", "NA unless all nine components are non-missing", "Strict missing propagation: do not treat NA as zero.", "Expense level and later financial indicators", "Formula: sum of nine components only when ExpenseComplete=1."],
        "NetBalance": ["NetBalance", "TotalIncome minus TotalExpense", "TotalIncome - TotalExpense", "Numeric", "Ratio", "Derived financial indicator", "NA unless both totals are available", "Missing totals propagate to NA.", "Financial position", "No partial totals used."],
        "ExpenseRatio": ["ExpenseRatio", "TotalExpense divided by TotalIncome × 100", "TotalExpense / TotalIncome × 100", "Numeric", "Ratio", "Derived financial indicator", "NA if either total is NA or TotalIncome=0", "No division by zero; missing totals propagate.", "Spending relative to income", "Stored as percent points, not Excel fraction."],
        "IncomeComplete": ["IncomeComplete", "All four income components are available", "Income component completeness", "Integer", "Nominal", "QA flag", "0=No, 1=Yes", "No imputation; transparent completeness flag.", "Sample-size reporting", "Required for valid TotalIncome."],
        "ExpenseComplete": ["ExpenseComplete", "All nine expense components are available", "Expense component completeness", "Integer", "Nominal", "QA flag", "0=No, 1=Yes", "No imputation; transparent completeness flag.", "Sample-size reporting", "Required for valid TotalExpense."],
        "FinancialIndicatorsAvailable": ["FinancialIndicatorsAvailable", "TotalIncome and TotalExpense are available", "TotalIncome and TotalExpense", "Integer", "Nominal", "QA flag", "0=No, 1=Yes", "Missing totals propagate.", "Sample-size reporting", "Required for NetBalance and ExpenseRatio."],
    }
    for ratio, expense in zip(RATIOS, [x.replace("_Clean", "") for x in EXPENSE]):
        dictionary_rows[ratio] = [ratio, f"{expense} divided by TotalExpense × 100", f"{expense} / TotalExpense × 100", "Numeric", "Ratio", "Expense composition", "NA if category or TotalExpense is NA, or TotalExpense=0", "Strict denominator and missing-value rules.", "Expense composition analysis", "Stored as percent points; eligible ratios sum to approximately 100."]
    for name, values in dictionary_rows.items():
        if name in entries:
            for col, value in enumerate(values, 1): dws.cell(entries[name], col).value = value
        else:
            dws.append(values)
    style_data_sheet(dws)
    dws.column_dimensions["B"].width = 36; dws.column_dimensions["C"].width = 48; dws.column_dimensions["H"].width = 48; dws.column_dimensions["J"].width = 46
    db.save(DICTIONARY)

    # Final in-memory/data assertions.
    assert len(rows) == 30
    assert sha256(RAW) == SOURCE_HASH
    assert [["" if x is None else str(x) for x in row] for row in sheet_matrix(load_workbook(MAIN, read_only=True)["Raw_Copy"])] == raw_matrix
    assert all((r["TotalIncome"] is not None) == bool(r["IncomeComplete"]) for r in rows)
    assert all((r["TotalExpense"] is not None) == bool(r["ExpenseComplete"]) for r in rows)
    for r in rows:
        if r["IncomeComplete"]: assert abs(r["TotalIncome"] - sum(r[x] for x in ["FamilyIncome", "PartTimeIncome", "ScholarshipIncome", "OtherIncome"])) < 1e-9
        if r["ExpenseComplete"]: assert abs(r["TotalExpense"] - sum(r[x] for x in ["FoodExpense", "TransportExpense", "HousingExpense", "UtilityExpense", "EducationExpense", "EntertainmentExpense", "ShoppingExpense", "SubscriptionExpense", "OtherExpense"])) < 1e-9
        if r["FinancialIndicatorsAvailable"]:
            assert abs(r["NetBalance"] - (r["TotalIncome"] - r["TotalExpense"])) < 1e-9
            assert abs(r["ExpenseRatio"] - r["TotalExpense"] / r["TotalIncome"] * 100) < 1e-9
        if r["TotalExpense"] not in (None, 0): assert abs(sum(r[x] for x in RATIOS) - 100) < 1e-8
    assert all(r["MoneySufficiency_Code"] in {1, 2, 3, 4} and r["MoneySufficiency_Group"] in {"Sufficient", "Insufficient"} and r["ShortageMonths_Code"] in {0,1,2,3} for r in rows)
    assert not any(np.isinf(x) for r in rows for x in r.values() if isinstance(x, (int, float)))
    sample = {"demographics":30, "TotalIncome":sum(r["TotalIncome"] is not None for r in rows), "TotalExpense":sum(r["TotalExpense"] is not None for r in rows), "NetBalance":sum(r["NetBalance"] is not None for r in rows), "ExpenseRatio":sum(r["ExpenseRatio"] is not None for r in rows), "CategoryRatios":sum(r["TotalExpense"] not in (None,0) for r in rows), "SavingAmount":sum(r["SavingAmount"] is not None for r in rows), "MoneySufficiency":sum(r["MoneySufficiency_Code"] is not None for r in rows), "ShortageMonths":sum(r["ShortageMonths_Code"] is not None for r in rows)}
    sample.update({x: sum(r[x] is not None for r in rows) for x in ["PlanSpending_Code", "TrackExpenses_Code", "ControlExpenses_Code", "PrioritizeNeeds_Code", "ImpulsePurchase_Code"]})
    summary = {"analysis_dimensions":[len(rows),len(headers)], "income_complete":sample["TotalIncome"], "expense_complete":sample["TotalExpense"], "financial_indicators":sample["NetBalance"], "sample_sizes":sample, "missing_total_income":30-sample["TotalIncome"], "missing_total_expense":30-sample["TotalExpense"], "housing_32000":"retained as reported"}
    (ROOT / "analysis_data_run_summary.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")
    print(json.dumps(summary, indent=2))

if __name__ == "__main__": main()
