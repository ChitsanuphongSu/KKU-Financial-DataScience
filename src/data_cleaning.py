"""Reproducible, non-destructive cleaning workflow for the KKU financial survey.

RAW_DATA.csv is read only.  This script preserves the original workbook sheet and
creates audited copies/working outputs without deriving any financial indicators.
"""
from __future__ import annotations

import csv
import hashlib
import json
import re
from collections import Counter
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "RAW_DATA.csv"
WORKBOOK = ROOT / "KKU_Financial_Cleaning.xlsx"
DICTIONARY = ROOT / "Data_Dictionary.xlsx"
LOGBOOK = ROOT / "Cleaning_Log.xlsx"

NUMERIC = {
    "FamilyIncome": 7, "PartTimeIncome": 8, "ScholarshipIncome": 9,
    "OtherIncome": 10, "FoodExpense": 12, "TransportExpense": 13,
    "HousingExpense": 14, "UtilityExpense": 15, "EducationExpense": 16,
    "EntertainmentExpense": 17, "ShoppingExpense": 18,
    "SubscriptionExpense": 19, "OtherExpense": 20, "SavingAmount": 23,
}
MAPPING = {
    "Timestamp": 0, "Consent": 1, "KKUStudent": 2, "Year": 3, "Faculty": 4,
    "Accommodation": 5, "PartTime": 6, **NUMERIC,
    "OtherIncomeSource": 11, "OtherExpenseDescription": 21, "SavingStatus": 22,
    "SavingPattern": 24, "PlanSpending": 25, "TrackExpenses": 26,
    "ControlExpenses": 27, "PrioritizeNeeds": 28, "ImpulsePurchase": 29,
    "MoneySufficiency": 30, "ShortageMonths": 31,
    "FinancialFactorsComment": 32,
}
EXPENSES = ["FoodExpense", "TransportExpense", "HousingExpense", "UtilityExpense",
            "EducationExpense", "EntertainmentExpense", "ShoppingExpense",
            "SubscriptionExpense", "OtherExpense"]
INCOMES = ["FamilyIncome", "PartTimeIncome", "ScholarshipIncome", "OtherIncome"]
LIKERT = ["PlanSpending", "TrackExpenses", "ControlExpenses", "PrioritizeNeeds", "ImpulsePurchase"]
MONEY_MAP = {
    "Level 1 — เพียงพอและมีเงินเหลือ": 1, "Level 2 — เพียงพอพอดี": 2,
    "Level 3 — ไม่เพียงพอบางเดือน": 3, "Level 4 — ไม่เพียงพอเป็นประจำ": 4,
}
SHORTAGE_RE = re.compile(r"^([0-3])\s*เดือน$")
EXACT_MONEY_RE = re.compile(r"^\s*(\d+(?:\.\d+)?)\s*(?:บาท)?\s*$")

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(bold=True, size=14)


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def read_source() -> tuple[list[str], list[list[str]]]:
    with RAW.open("r", encoding="utf-8-sig", newline="") as f:
        matrix = list(csv.reader(f))
    return matrix[0], matrix[1:]


def parse_exact_money(value: str) -> tuple[float | None, str | None, str | None]:
    """Return (numeric value, issue type, reason); never infer ambiguous amounts."""
    raw = value.strip()
    if raw == "":
        return None, "missing_value", "Blank response; retained as missing."
    if "+" in raw:
        return None, "lower_bound", "Lower-bound response is not an exact amount."
    if "/" in raw:
        return None, "period_ambiguous", "Period is ambiguous; no monthly conversion assumed."
    candidate = raw.replace(",", "")
    match = EXACT_MONEY_RE.match(candidate)
    if not match:
        return None, "invalid_numeric", "Text or non-numeric response cannot be reliably converted."
    number = float(match.group(1))
    if number.is_integer():
        number = int(number)
    canonical = str(number)
    if raw != canonical:
        return number, "format_standardization", "Unambiguous numeric formatting standardized."
    return number, None, None


def setup_sheet(ws, widths=True):
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    if ws.max_row:
        for cell in ws[1]:
            cell.fill, cell.font = HEADER_FILL, HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if widths:
        for col in range(1, ws.max_column + 1):
            values = [str(ws.cell(r, col).value or "") for r in range(1, min(ws.max_row, 50) + 1)]
            ws.column_dimensions[get_column_letter(col)].width = min(max(max(map(len, values), default=10) + 2, 12), 32)


def replace_sheet(wb, name):
    if name in wb.sheetnames:
        del wb[name]
    return wb.create_sheet(name)


def append_matrix(ws, rows):
    for row in rows:
        ws.append(list(row))


def main():
    before_hash = sha256(RAW)
    headers, rows = read_source()
    assert len(headers) == 33 and len(rows) == 30, "Unexpected source dimensions."
    assert all(len(row) == len(headers) for row in rows), "Ragged CSV rows detected."
    source = pd.DataFrame(rows, columns=headers)
    logs: list[dict] = []
    clean_rows: list[dict] = []
    numeric_modified = numeric_na = standardized = 0

    for row_num, raw in enumerate(rows, start=2):
        record = {"Row_ID": row_num - 1}
        for var, idx in MAPPING.items():
            record[var] = raw[idx]
        for var, idx in NUMERIC.items():
            value, issue, reason = parse_exact_money(raw[idx])
            record[f"{var}_Clean"] = value
            if issue:
                action = "Standardized exact numeric format" if issue == "format_standardization" else "Set cleaned value to NA"
                if issue == "format_standardization":
                    standardized += 1
                else:
                    numeric_na += 1
                numeric_modified += 1
                logs.append({"Row": row_num, "Variable": var, "Original_Value": raw[idx],
                             "Cleaned_Value": "" if value is None else value, "Issue_Type": issue,
                             "Action": action, "Reason": reason, "Status": "FIXED" if value is not None else "REQUIRES REVIEW",
                             "Reviewer_Note": ""})
        for var in LIKERT:
            raw_value = raw[MAPPING[var]].strip()
            coded = int(raw_value) if raw_value in {"1", "2", "3", "4", "5"} else None
            record[f"{var}_Code"] = coded
            if coded is None:
                logs.append({"Row": row_num, "Variable": var, "Original_Value": raw_value, "Cleaned_Value": "",
                             "Issue_Type": "category_inconsistency", "Action": "Set coded value to NA",
                             "Reason": "Not one of observed valid Likert responses 1–5.", "Status": "REQUIRES REVIEW", "Reviewer_Note": ""})
        money = raw[MAPPING["MoneySufficiency"]].strip()
        record["MoneySufficiency_Code"] = MONEY_MAP.get(money)
        record["MoneySufficiency_Group"] = ("Sufficient" if MONEY_MAP.get(money) in (1, 2) else "Insufficient"
                                             if MONEY_MAP.get(money) in (3, 4) else "")
        if money not in MONEY_MAP:
            logs.append({"Row": row_num, "Variable": "MoneySufficiency", "Original_Value": money, "Cleaned_Value": "",
                         "Issue_Type": "category_inconsistency", "Action": "Set coded value to NA",
                         "Reason": "Response does not match the observed four-level questionnaire scale.",
                         "Status": "REQUIRES REVIEW", "Reviewer_Note": ""})
        shortage = raw[MAPPING["ShortageMonths"]].strip()
        m = SHORTAGE_RE.match(shortage)
        record["ShortageMonths_Code"] = int(m.group(1)) if m else None
        if not m:
            logs.append({"Row": row_num, "Variable": "ShortageMonths", "Original_Value": shortage, "Cleaned_Value": "",
                         "Issue_Type": "out_of_range", "Action": "Set coded value to NA",
                         "Reason": "Expected 0–3 months in the stated response format.",
                         "Status": "REQUIRES REVIEW", "Reviewer_Note": ""})
        clean_rows.append(record)

    # IQR screening only; no values are changed.
    outliers = []
    for var in NUMERIC:
        vals = pd.Series([r[f"{var}_Clean"] for r in clean_rows], dtype="float64").dropna()
        if len(vals) >= 4:
            q1, q3 = vals.quantile(.25), vals.quantile(.75)
            lo, hi = q1 - 1.5 * (q3-q1), q3 + 1.5 * (q3-q1)
            for i, value in enumerate(vals):
                if value < lo or value > hi:
                    # Index returned by pandas is record index.
                    respondent = vals.index[i] + 2
                    kind = "possible_data_entry_error" if (var == "HousingExpense" and value == 32000) else "plausible_extreme"
                    outliers.append([respondent, var, value, round(lo, 2), round(hi, 2), kind, "REQUIRES REVIEW" if kind.startswith("possible") else "ACCEPTED AS VALID"])
                    logs.append({"Row": respondent, "Variable": var, "Original_Value": rows[respondent-2][NUMERIC[var]],
                                 "Cleaned_Value": value, "Issue_Type": "possible_outlier", "Action": "No value changed",
                                 "Reason": f"IQR screen bounds: {lo:.2f} to {hi:.2f}.", "Status": "REQUIRES REVIEW" if kind.startswith("possible") else "ACCEPTED AS VALID", "Reviewer_Note": kind})

    # Add response-level semantic checks without editing data.
    for i, r in enumerate(clean_rows, start=2):
        if r["PartTime"] == "ไม่มี" and r["PartTimeIncome_Clean"] not in (0, None):
            logs.append({"Row": i, "Variable": "PartTimeIncome", "Original_Value": rows[i-2][8], "Cleaned_Value": r["PartTimeIncome_Clean"],
                         "Issue_Type": "category_inconsistency", "Action": "No value changed", "Reason": "Part-time status says no, but income is non-zero.", "Status": "REQUIRES REVIEW", "Reviewer_Note": "Confirm questionnaire interpretation."})

    # Preserve the supplied workspace's original sheet(s) and append requested sheets.
    wb = load_workbook(WORKBOOK)
    raw_ws = replace_sheet(wb, "Raw_Copy")
    append_matrix(raw_ws, [headers] + rows)
    setup_sheet(raw_ws)
    raw_ws.auto_filter.ref = raw_ws.dimensions

    readme = replace_sheet(wb, "README")
    readme.append(["KKU Financial Survey: cleaning workspace"])
    readme["A1"].font = TITLE_FONT
    for line in [
        ["Source", "RAW_DATA.csv (immutable; SHA-256 recorded at execution)"],
        ["Scope", "Cleaning and validation only. No EDA, correlations, group comparisons, or derived financial indicators."],
        ["Raw_Copy", "Exact cell-level CSV copy; all cleaning occurs in Working_Data."],
        ["Numeric rule", "Only unambiguous numbers (optional comma or บาท) are standardized. Lower bounds, periods, and text become NA only in cleaned fields."],
        ["Derived indicators", "Headers are prepared but deliberately blank until component completeness policy is approved."],
        ["Reproducibility", "Run src/data_cleaning.py from this project folder."],
        ["Source SHA-256", before_hash],
    ]: readme.append(line)
    readme.column_dimensions["A"].width, readme.column_dimensions["B"].width = 24, 110
    readme.sheet_view.showGridLines = False

    working = replace_sheet(wb, "Working_Data")
    original_vars = ["Row_ID"] + list(MAPPING.keys())
    clean_vars = [f"{v}_Clean" for v in NUMERIC]
    code_vars = [f"{v}_Code" for v in LIKERT] + ["MoneySufficiency_Code", "MoneySufficiency_Group", "ShortageMonths_Code"]
    derived_vars = ["TotalIncome", "TotalExpense", "NetBalance", "ExpenseRatio", "CategoryRatio"]
    working_headers = original_vars + clean_vars + code_vars + derived_vars
    working.append(working_headers)
    for r in clean_rows:
        working.append([r.get(c, None) for c in working_headers])
    setup_sheet(working)
    working.auto_filter.ref = working.dimensions
    for c in range(1, working.max_column + 1):
        if working.cell(1, c).value in derived_vars:
            working.cell(1, c).fill = PatternFill("solid", fgColor="808080")
    for row in range(2, working.max_row + 1):
        for c in range(1, working.max_column + 1):
            if working.cell(1, c).value in derived_vars:
                working.cell(row, c).value = None

    mapping_ws = replace_sheet(wb, "Variable_Mapping")
    mapping_ws.append(["Variable", "Source_Column_Number", "Exact_Source_Column", "Status"])
    for variable, idx in MAPPING.items():
        mapping_ws.append([variable, idx + 1, headers[idx], "Mapped from actual CSV header"])
    setup_sheet(mapping_ws)

    report = replace_sheet(wb, "Data_Quality_Report")
    report.append(["Data Quality Report"]); report["A1"].font = TITLE_FONT
    report.append(["Metric", "Value", "Classification", "Notes"])
    for c in report[2]: c.fill, c.font = HEADER_FILL, HEADER_FONT
    unresolved = sum(x["Status"] == "REQUIRES REVIEW" for x in logs)
    report_rows = [
        ["Total respondents", len(rows), "ACCEPTED AS VALID", "All 30 retained."],
        ["Total source variables", len(headers), "ACCEPTED AS VALID", "CSV source dimensions: 30 × 33."],
        ["Exact duplicate responses", int(source.duplicated().sum()), "ACCEPTED AS VALID", "No full-row duplicates."],
        ["Numeric fields inspected", len(NUMERIC), "ACCEPTED AS VALID", "Four income, nine expense, and one saving field."],
        ["Standardized numeric cells", standardized, "FIXED", "Only unambiguous comma/บาท formatting."],
        ["Numeric cells converted to NA", numeric_na, "FIXED / REQUIRES REVIEW", "Text, lower bounds, and ambiguous periods were not guessed."],
        ["Potential outliers", len(outliers), "REQUIRES REVIEW / ACCEPTED AS VALID", "IQR screen only; none removed or capped."],
        ["Unresolved log entries", unresolved, "REQUIRES REVIEW", "Includes invalid/ambiguous cells and possible data-entry errors."],
        ["Derived indicators", "Not calculated", "EXCLUDED ONLY FOR SPECIFIC ANALYSIS", "Missing components are not assumed to be zero."],
        ["Likert scale", "1–5 observed", "ACCEPTED AS VALID", "All five behavior fields are numeric categories 1–5."],
        ["Money sufficiency", "4 observed levels", "ACCEPTED AS VALID", "Coded separately; original text retained."],
        ["Shortage months", "0–3 observed", "ACCEPTED AS VALID", "Parsed from 0/1/2/3 เดือน without silent corrections."],
    ]
    append_matrix(report, report_rows)
    report.append([]); report.append(["Missing Values by Source Variable"])
    report.append(["Source_Column_Number", "Source_Column", "Blank_Count", "Blank_Percent"])
    for idx, header in enumerate(headers, start=1):
        blanks = sum(not row[idx-1].strip() for row in rows)
        report.append([idx, header, blanks, blanks / len(rows)])
    for cell in report[report.max_row - len(headers):report.max_row + 1]:
        if len(cell) >= 4:
            cell[3].number_format = "0.0%"
    report.append([]); report.append(["Numeric Cleaning by Variable"])
    report.append(["Variable", "Invalid_or_Ambiguous_to_NA", "Format_Standardized", "Cleaned_Missing_Count"])
    for variable in NUMERIC:
        relevant = [x for x in logs if x["Variable"] == variable]
        na_count = sum(x["Issue_Type"] in {"invalid_numeric", "lower_bound", "period_ambiguous", "missing_value"} for x in relevant)
        standardized_count = sum(x["Issue_Type"] == "format_standardization" for x in relevant)
        report.append([variable, na_count, standardized_count, na_count])
    report.append([]); report.append(["Potential Outliers (IQR screen only)"])
    report.append(["Row", "Variable", "Value", "Lower_Bound", "Upper_Bound", "Assessment", "Status"])
    append_matrix(report, outliers or [["None", "", "", "", "", "", ""]])
    for col in range(1, 8): report.column_dimensions[get_column_letter(col)].width = 26
    report.sheet_view.showGridLines = False

    # Write log inside the main workbook and as a dedicated deliverable.
    log_cols = ["Log_ID", "Row", "Variable", "Original_Value", "Cleaned_Value", "Issue_Type", "Action", "Reason", "Status", "Reviewer_Note"]
    log_rows = []
    for number, entry in enumerate(logs, 1):
        log_rows.append([number] + [entry.get(c, "") for c in log_cols[1:]])
    log_ws = replace_sheet(wb, "Cleaning_Log")
    append_matrix(log_ws, [log_cols] + log_rows)
    setup_sheet(log_ws)
    log_ws.auto_filter.ref = log_ws.dimensions

    wb.save(WORKBOOK)
    # Dedicated Cleaning_Log workbook
    logbook = Workbook(); lws = logbook.active; lws.title = "Cleaning_Log"
    append_matrix(lws, [log_cols] + log_rows); setup_sheet(lws); lws.auto_filter.ref = lws.dimensions
    logbook.save(LOGBOOK)

    # Data dictionary, matching the implemented Working_Data fields and source columns.
    dictionary = Workbook(); dws = dictionary.active; dws.title = "Data_Dictionary"
    dict_cols = ["Variable", "Description", "Source_Column", "Type", "Measurement_Level", "Role", "Valid_Values_or_Range", "Cleaning_Guidance", "Analysis_Use", "Notes"]
    dws.append(dict_cols)
    descriptions = {
        "Year":"Student year", "Faculty":"Faculty/college", "Accommodation":"Current accommodation", "PartTime":"Has part-time work",
        "MoneySufficiency":"Monthly money sufficiency response", "ShortageMonths":"Months insufficient in prior 3 months",
        "FinancialFactorsComment":"Open-ended perceived financial factors", "SavingStatus":"Saved in past month", "SavingPattern":"Usual saving pattern",
    }
    for var, idx in MAPPING.items():
        if var in NUMERIC: typ, level, valid, guidance, use = "Numeric", "Ratio", "0 or positive exact amount", "Unambiguous numeric text only; otherwise cleaned field is NA.", "Input to later indicators after validation"
        elif var in LIKERT: typ, level, valid, guidance, use = "Categorical code", "Ordinal", "Observed 1–5", "Preserve raw response; code only observed categories.", "Financial behavior comparison"
        elif var == "MoneySufficiency": typ, level, valid, guidance, use = "Categorical", "Ordinal", "Four observed Level 1–4 labels", "Original text retained; separately code 1–4.", "Main outcome"
        elif var == "ShortageMonths": typ, level, valid, guidance, use = "Categorical count", "Ordinal", "0–3 months", "Parse only stated valid values; flag others.", "Supporting outcome"
        elif var in {"FinancialFactorsComment", "OtherIncomeSource", "OtherExpenseDescription"}: typ, level, valid, guidance, use = "Text", "Text / N/A", "Free text / optional", "Retain as supplied.", "Context only"
        elif var == "Timestamp": typ, level, valid, guidance, use = "Datetime", "Interval", "Survey submission timestamp", "Retain source text.", "Provenance only"
        else: typ, level, valid, guidance, use = "Categorical", "Nominal", "Observed source categories", "Retain source response.", "Grouping/context"
        dws.append([var, descriptions.get(var, var), headers[idx], typ, level, "Source / working variable", valid, guidance, use, "Original response retained in Raw_Copy."])
    for var in NUMERIC: dws.append([f"{var}_Clean", f"Validated exact numeric version of {var}", headers[NUMERIC[var]], "Numeric", "Ratio", "Cleaned working variable", "0 or positive exact amount; NA if unresolved", "No guessing; zeros retained.", "Future financial indicators", "Derived indicators intentionally not yet calculated."])
    for var in LIKERT: dws.append([f"{var}_Code", f"Coded {var}", headers[MAPPING[var]], "Integer", "Ordinal", "Cleaned coded variable", "1–5", "Observed numeric scale used unchanged.", "Financial behavior comparison", "Higher ImpulsePurchase = more frequent only if questionnaire direction confirms it."])
    dws.append(["MoneySufficiency_Code", "Coded money sufficiency", headers[30], "Integer", "Ordinal", "Main outcome coded", "1=surplus; 2=just enough; 3=insufficient some months; 4=insufficient regularly", "Separate from raw label.", "Future group comparison", "Sufficient=1–2; Insufficient=3–4."])
    dws.append(["MoneySufficiency_Group", "Future group-comparison grouping", headers[30], "Categorical", "Nominal", "Derived grouping", "Sufficient / Insufficient", "Derived from valid code only.", "Future group comparison", "No comparison performed."])
    dws.append(["ShortageMonths_Code", "Parsed shortage-month count", headers[31], "Integer", "Ordinal", "Supporting outcome coded", "0–3", "Only explicit 0–3 เดือน values.", "Supporting outcome", "Original label retained."])
    for var in ["TotalIncome", "TotalExpense", "NetBalance", "ExpenseRatio", "CategoryRatio"]:
        dws.append([var, "Prepared only; deliberately blank", "N/A", "Numeric", "Ratio", "Future derived indicator", "Not calculated", "Do not replace missing components with zero.", "Future analysis", "Requires approved component completeness rule."])
    setup_sheet(dws); dws.auto_filter.ref = dws.dimensions
    dws.column_dimensions["B"].width = 36; dws.column_dimensions["C"].width = 48; dws.column_dimensions["H"].width = 48; dws.column_dimensions["J"].width = 46
    dictionary.save(DICTIONARY)

    after_hash = sha256(RAW)
    assert before_hash == after_hash, "RAW_DATA.csv changed unexpectedly."
    check_wb = load_workbook(WORKBOOK, read_only=True, data_only=False)
    assert check_wb["Raw_Copy"].max_row == len(rows) + 1
    assert check_wb["Raw_Copy"].max_column == len(headers)
    assert check_wb["Working_Data"].max_row == len(rows) + 1
    assert all(r[f"{v}_Clean"] == 0 for r in clean_rows for v in NUMERIC if rows[r["Row_ID"]-1][NUMERIC[v]].strip() == "0")
    assert next(r for r in clean_rows if r["Row_ID"] == 16)["ShoppingExpense_Clean"] is None, "Lower-bound 5000+ must remain unresolved."
    assert next(r for r in clean_rows if r["Row_ID"] == 21)["HousingExpense_Clean"] is None, "Per-term amount must remain unresolved."
    assert set(x["MoneySufficiency_Code"] for x in clean_rows) == {1, 2, 3, 4}
    assert all(x["ShortageMonths_Code"] in {0,1,2,3} for x in clean_rows)
    assert all(x.get(v) is None for x in clean_rows for v in ["TotalIncome", "TotalExpense", "NetBalance", "ExpenseRatio", "CategoryRatio"])
    summary = {"source_dimensions": [len(rows), len(headers)], "respondents_retained": len(clean_rows),
               "cells_modified": numeric_modified, "cells_converted_to_na": numeric_na, "standardized_numeric": standardized,
               "unresolved_issues": unresolved, "outliers": outliers, "log_entries": len(logs), "raw_sha256": before_hash}
    encoded_summary = json.dumps(summary, ensure_ascii=False, indent=2, default=lambda x: x.item() if hasattr(x, "item") else str(x))
    (ROOT / "cleaning_run_summary.json").write_text(encoded_summary, encoding="utf-8")
    print(encoded_summary)


if __name__ == "__main__":
    main()
